"""One-off diff between LLM-produced PPTX and Bernardo's hand-revised version."""
from pptx import Presentation
from pptx.util import Emu
import sys

def emu_to_inches(emu):
    return round(emu / 914400, 2) if emu is not None else None

def dump_slide(slide, idx):
    out = [f"\n===== SLIDE {idx} ====="]
    out.append(f"  shapes: {len(slide.shapes)}")
    for s_idx, shape in enumerate(slide.shapes):
        try:
            left = emu_to_inches(shape.left)
            top = emu_to_inches(shape.top)
            w = emu_to_inches(shape.width)
            h = emu_to_inches(shape.height)
        except Exception:
            left = top = w = h = None
        kind = shape.shape_type
        name = shape.name
        out.append(f"  [{s_idx}] {kind} name={name!r} pos=({left},{top}) size=({w}x{h})")
        if shape.has_text_frame:
            for p_idx, para in enumerate(shape.text_frame.paragraphs):
                runs_info = []
                for r in para.runs:
                    sz = r.font.size.pt if r.font.size else None
                    bold = r.font.bold
                    color = None
                    try:
                        if r.font.color and r.font.color.rgb:
                            color = str(r.font.color.rgb)
                    except Exception:
                        pass
                    fname = r.font.name
                    runs_info.append(f"sz={sz} b={bold} c={color} fnt={fname} txt={r.text!r}")
                out.append(f"    P{p_idx}: " + " | ".join(runs_info) if runs_info else f"    P{p_idx}: (empty)")
        if shape.has_table:
            tbl = shape.table
            out.append(f"    TABLE {len(tbl.rows)}x{len(tbl.columns)}")
            for r_idx, row in enumerate(tbl.rows):
                cells = []
                for cell in row.cells:
                    cells.append(cell.text_frame.text[:40].replace("\n", " "))
                out.append(f"      R{r_idx}: " + " | ".join(cells))
    return "\n".join(out)

def dump_pptx(path, label):
    prs = Presentation(path)
    out = [f"### {label} · {path}"]
    out.append(f"slide_w={emu_to_inches(prs.slide_width)} slide_h={emu_to_inches(prs.slide_height)}")
    out.append(f"n_slides={len(prs.slides)}")
    for i, slide in enumerate(prs.slides, 1):
        out.append(dump_slide(slide, i))
    return "\n".join(out)

if __name__ == "__main__":
    base = r"CONTENT\PARTE2-Materiales\II1-Ingesta_Canal"
    llm = f"{base}\\M1_slides_II1-v6-llm0.pptx"
    rev = f"{base}\\M1_slides_II1-v6.pptx"
    with open("sessions/_diff_llm.txt", "w", encoding="utf-8") as f:
        f.write(dump_pptx(llm, "LLM"))
    with open("sessions/_diff_rev.txt", "w", encoding="utf-8") as f:
        f.write(dump_pptx(rev, "REVISED"))
    print("wrote sessions/_diff_llm.txt and sessions/_diff_rev.txt")
