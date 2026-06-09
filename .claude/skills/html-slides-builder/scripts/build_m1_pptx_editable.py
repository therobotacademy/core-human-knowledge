#!/usr/bin/env python3
"""
build_m1_pptx_editable.py — APM M1 HTML -> editable PPTX

Parses each slide section from the APM HTML deck and builds python-pptx
shapes (text boxes, filled rectangles, images) with proper colors, fonts,
and absolute positioning. All text is fully editable in PowerPoint.

Requirements:
    pip install python-pptx beautifulsoup4 pillow lxml
    Fonts: install Syne + IBM Plex Sans + IBM Plex Mono from Google Fonts
           for correct rendering in PowerPoint.

Usage:
    python scripts/build_m1_pptx_editable.py CONTENT/PARTE2-Materiales/B1-Concepto_Agente/M1_slides_B1.html

    # All sessions (PowerShell):
    foreach ($f in Get-ChildItem "CONTENT/PARTE2-Materiales/B*/M1_slides_B*.html") {
        python scripts/build_m1_pptx_editable.py $f.FullName
    }

Output: same directory, basename + _editable.pptx
"""

import sys
import re
from pathlib import Path
from urllib.parse import unquote

from bs4 import BeautifulSoup, NavigableString, Tag
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette ────────────────────────────────────────────────────────────────────
AZUL        = RGBColor(0x1E, 0x3A, 0x5F)
AZUL_MEDIO  = RGBColor(0x2E, 0x6B, 0x9E)
CREMA       = RGBColor(0xFA, 0xFA, 0xF7)
NARANJA_LUZ = RGBColor(0xFF, 0x8C, 0x3B)
NARANJA_OSC = RGBColor(0xC2, 0x51, 0x0A)
VERDE       = RGBColor(0x0D, 0x7C, 0x5A)
SEC         = RGBColor(0x6B, 0x6B, 0x6B)
TEXTO       = RGBColor(0x1C, 0x1C, 0x1C)
BLUE_WHITE  = RGBColor(0x9A, 0xB0, 0xC8)
DIM         = RGBColor(0x9B, 0x9B, 0x9B)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BLACK       = RGBColor(0x00, 0x00, 0x00)
CARD_BG     = RGBColor(0xFF, 0xFF, 0xFF)
BORDE_C     = RGBColor(0xE8, 0xE3, 0xD8)

SYNE  = "Syne"
IBM_S = "IBM Plex Sans"
IBM_M = "IBM Plex Mono"

# ── Dimensions: HTML 960x540 -> PPTX 10"x5.625" ───────────────────────────────
W = Inches(10)
H = Inches(5.625)
_EMU_PER_PX = round(W / 960)  # 9525

def px(n):
    """HTML pixel -> EMU."""
    return round(n * _EMU_PER_PX)

PAD_L  = px(80)
PAD_T  = px(60)
HDR_H  = px(50)
BODY_L = px(48)
BODY_T = px(32)
BOX_W  = W - PAD_L - PAD_L  # 800px wide content area

# ── CSS helpers ────────────────────────────────────────────────────────────────
_CSS_VARS = {
    '--azul': AZUL, '--azul-medio': AZUL_MEDIO, '--crema': CREMA,
    '--naranja-luz': NARANJA_LUZ, '--naranja-oscuro': NARANJA_OSC,
    '--verde': VERDE, '--sec': SEC, '--texto': TEXTO,
    '--blue-white': BLUE_WHITE, '--dim': DIM,
}

def resolve_color(val):
    if not val:
        return None
    val = val.strip()
    m = re.match(r'var\((--[\w-]+)\)', val)
    if m:
        return _CSS_VARS.get(m.group(1))
    if val in ('#fff', '#ffffff', 'white'):
        return WHITE
    if val.startswith('#'):
        h = val[1:]
        if len(h) == 3:
            h = h[0]*2 + h[1]*2 + h[2]*2
        return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))
    return None

def parse_inline_style(s):
    """'font-size:26px; color:red' -> {'font-size':'26px','color':'red'}"""
    d = {}
    for part in (s or '').split(';'):
        if ':' in part:
            k, v = part.split(':', 1)
            d[k.strip()] = v.strip()
    return d

def css_px_to_pt(val):
    m = re.search(r'(\d+(?:\.\d+)?)px', val or '')
    return Pt(round(float(m.group(1)) * 0.75, 1)) if m else None

def css_px_to_emu(val):
    m = re.search(r'(\d+(?:\.\d+)?)px', val or '')
    return px(float(m.group(1))) if m else 0

# ── PPTX shape helpers ─────────────────────────────────────────────────────────

def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, l, t, w, h, fill=None, line_color=None, line_width=Pt(0.75)):
    shape = slide.shapes.add_shape(1, l, t, w, h)  # 1 = RECTANGLE
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape

def add_tb(slide, l, t, w, h):
    """Add textbox with zero internal margins."""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf

def _first_para(tf):
    """Return first paragraph (always exists in a new textbox)."""
    return tf.paragraphs[0]

def set_run(run, text, font=IBM_S, size=Pt(11), color=TEXTO,
            bold=False, italic=False):
    run.text = text
    run.font.name = font
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic

def simple_para(tf, text, font=IBM_S, size=Pt(11), color=TEXTO,
                bold=False, italic=False, align=PP_ALIGN.LEFT,
                space_before=Pt(0), space_after=Pt(4)):
    """Write one paragraph of plain text into tf (uses first para or appends)."""
    if tf.paragraphs[0].text == '':
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = align
    p.space_before = space_before
    p.space_after = space_after
    if text:
        r = p.add_run()
        set_run(r, text, font, size, color, bold, italic)
    return p

def add_counter(slide, idx, total):
    tf = add_tb(slide, W - px(70), H - px(30), px(60), px(22))
    simple_para(tf, f"{idx} / {total}", IBM_M, Pt(7.5), DIM, align=PP_ALIGN.RIGHT)

# ── Rich-text inline parser ────────────────────────────────────────────────────

def parse_inline_node(node, para,
                      font=IBM_S, size=Pt(11), color=TEXTO,
                      bold=False, italic=False):
    """
    Recursively add runs to `para` from an HTML node.
    Handles NavigableString, <strong>, <em>, <span style="...">.
    Caller handles <br> (splits into new paragraphs upstream).
    """
    if isinstance(node, NavigableString):
        text = str(node)
        if text.strip() or (text and para.runs):
            r = para.add_run()
            set_run(r, text, font, size, color, bold, italic)
        return
    if not isinstance(node, Tag):
        return

    tag = (node.name or '').lower()
    if tag == 'br':
        return  # handled at paragraph split level

    st = parse_inline_style(node.get('style', ''))
    f2 = font
    s2 = css_px_to_pt(st.get('font-size', '')) or size
    c2 = resolve_color(st.get('color')) or color
    b2 = bold or tag == 'strong'
    i2 = italic or tag in ('em', 'i')

    for child in node.children:
        parse_inline_node(child, para, f2, s2, c2, b2, i2)


def split_by_br(element):
    """
    Yield groups of sibling nodes separated by <br> tags.
    Each group = one line (list of nodes).
    """
    current = []
    for child in element.children:
        if isinstance(child, Tag) and child.name == 'br':
            yield current
            current = []
        else:
            current.append(child)
    if current:
        yield current

def nodes_text(nodes):
    return ''.join(
        str(n) if isinstance(n, NavigableString) else n.get_text()
        for n in nodes
    ).strip()

def add_multiline_element(tf, element,
                          font=SYNE, size=Pt(15), color=TEXTO,
                          align=PP_ALIGN.LEFT,
                          space_between=Pt(4)):
    """
    Parse an element with possible <br> tags, adding one paragraph per line.
    Returns tf for chaining.
    """
    first = True
    for line_nodes in split_by_br(element):
        text_content = nodes_text(line_nodes)
        if first:
            p = _first_para(tf)
            first = False
        else:
            p = tf.add_paragraph()

        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = space_between

        if not text_content:
            # empty line: paragraph acts as spacer
            continue

        for node in line_nodes:
            parse_inline_node(node, p, font, size, color)
    return tf

# ── Slide builders ─────────────────────────────────────────────────────────────

def build_cover(slide, sec, idx, total):
    set_slide_bg(slide, AZUL)
    add_rect(slide, 0, 0, px(5), H, fill=NARANJA_LUZ)

    y = PAD_T

    chip = sec.select_one('.chip')
    if chip:
        tf = add_tb(slide, PAD_L, y, BOX_W, px(22))
        simple_para(tf, chip.get_text(strip=True), IBM_M, Pt(7.5), NARANJA_LUZ)
        y += px(22 + 24)

    h1 = sec.find('h1')
    if h1:
        tf = add_tb(slide, PAD_L, y, BOX_W, px(130))
        # h1 may have <br>
        add_multiline_element(tf, h1, SYNE, Pt(42), WHITE, space_between=Pt(2))
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.bold = True
        y += px(130 + 20)

    sub = sec.select_one('.sub')
    if sub:
        tf = add_tb(slide, PAD_L, y, px(640), px(80))
        simple_para(tf, sub.get_text(strip=True), IBM_S, Pt(12), BLUE_WHITE,
                    space_after=Pt(0))

    footer = sec.select_one('.footer')
    if footer:
        tf = add_tb(slide, PAD_L, H - px(40), px(400), px(22))
        simple_para(tf, footer.get_text(strip=True), IBM_M, Pt(7.5), DIM)

    add_counter(slide, idx, total)


def build_idea(slide, sec, idx, total):
    set_slide_bg(slide, AZUL)
    add_rect(slide, 0, 0, px(5), H, fill=NARANJA_LUZ)

    y = PAD_T

    chip = sec.select_one('.chip')
    if chip:
        tf = add_tb(slide, PAD_L + px(20), y, BOX_W, px(22))
        simple_para(tf, chip.get_text(strip=True), IBM_M, Pt(7.5), NARANJA_LUZ)
        y += px(22 + 36)

    quote = sec.select_one('.quote')
    if quote:
        tf = add_tb(slide, PAD_L + px(20), y, px(760), px(180))
        add_multiline_element(tf, quote, SYNE, Pt(28.5), WHITE,
                              space_between=Pt(6))
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.bold = True
        y += px(180 + 28)

    fq = sec.select_one('.footer-quote')
    if fq:
        tf = add_tb(slide, PAD_L + px(20), y, px(760), px(36))
        simple_para(tf, fq.get_text(strip=True), IBM_S, Pt(10.5), BLUE_WHITE)

    add_counter(slide, idx, total)


def build_image_slide(slide, sec, html_dir, idx, total):
    set_slide_bg(slide, BLACK)

    img_tag = sec.find('img')
    if img_tag:
        src = unquote(img_tag.get('src', ''))
        # Try relative to html_dir first, then canonical imagenes/ path
        candidates = [
            html_dir / src,
            html_dir.parent / src,
        ]
        # Also try canonical storyboard imagenes/
        img_name = Path(src).name
        # html_dir is 3 levels deep: CONTENT/PARTE2-Materiales/Bx-*/
        repo_root = html_dir.parent.parent.parent
        candidates.append(
            repo_root / 'CONTENT' / 'PARTE1-Marco-comun' /
            'asimov-storyboard' / 'imagenes' / img_name
        )
        img_path = next((p for p in candidates if p.exists()), None)

        # Fuzzy fallback: match any M0 .png in imagenes/ sharing the same Bx code
        if img_path is None:
            session_match = re.search(r'B\d+', img_name)
            if session_match:
                code = session_match.group()
                imagenes_dir = (
                    repo_root / 'CONTENT' / 'PARTE1-Marco-comun' /
                    'asimov-storyboard' / 'imagenes'
                )
                matches = [
                    p for p in imagenes_dir.glob('*.png')
                    if code in p.name and 'M0' in p.name
                ]
                if matches:
                    img_path = matches[0]

        if img_path:
            pic = slide.shapes.add_picture(str(img_path), 0, 0, W, H)
        else:
            # Placeholder rect if image missing
            add_rect(slide, 0, 0, W, H, fill=RGBColor(0x33, 0x33, 0x33))
            tf = add_tb(slide, px(100), px(220), W - px(200), px(80))
            simple_para(tf, f"[imagen: {src}]", IBM_M, Pt(12), DIM,
                        align=PP_ALIGN.CENTER)

    add_counter(slide, idx, total)


def build_story(slide, sec, idx, total):
    """Story slides: crema bg, label, h2, lines, pivote."""
    set_slide_bg(slide, CREMA)

    y = PAD_T

    label = sec.select_one('.label')
    if label:
        tf = add_tb(slide, PAD_L, y, BOX_W, px(22))
        simple_para(tf, label.get_text(strip=True), IBM_M, Pt(6.75), SEC)
        y += px(22 + 28)

    for child in sec.children:
        if not isinstance(child, Tag):
            continue
        tag = child.name.lower()
        cls = child.get('class', [])
        st = parse_inline_style(child.get('style', ''))

        if tag == 'h2':
            lines_count = sum(1 for _ in split_by_br(child))
            h2_h = px(50 * max(1, lines_count))
            tf = add_tb(slide, PAD_L, y, px(760), h2_h)
            add_multiline_element(tf, child, SYNE, Pt(28.5), AZUL,
                                  space_between=Pt(4))
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.bold = True
            y += h2_h + px(16)

        elif tag == 'p' and 'label' not in cls:
            p_st = parse_inline_style(child.get('style', ''))
            col = resolve_color(p_st.get('color')) or TEXTO
            italic = 'italic' in p_st.get('font-style', '')
            sz = css_px_to_pt(p_st.get('font-size', '')) or Pt(12.75)
            mt = css_px_to_emu(p_st.get('margin-top', ''))
            y += mt
            tf = add_tb(slide, PAD_L, y, px(760), px(80))
            p = _first_para(tf)
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(4)
            for node in child.children:
                parse_inline_node(node, p, IBM_S, sz, col, italic=italic)
            y += px(60)

        elif tag == 'div' and 'lines' in cls:
            is_big = 'big' in cls
            sz = Pt(21) if is_big else Pt(15)
            align = PP_ALIGN.CENTER if 'text-align:center' in child.get('style', '') else PP_ALIGN.LEFT
            mt = css_px_to_emu(st.get('margin-top', ''))
            y += mt
            h_est = H - y - px(40)
            if align == PP_ALIGN.CENTER:
                tb_w = px(800)
                tb_l = (W - tb_w) // 2
            else:
                tb_w = BOX_W
                tb_l = PAD_L
            tf = add_tb(slide, tb_l, y, tb_w, h_est)
            add_multiline_element(tf, child, SYNE, sz, TEXTO, align=align,
                                  space_between=Pt(8) if is_big else Pt(4))
            y += h_est // 2  # advance roughly half the remaining space

        elif tag == 'div' and 'pivote' in cls:
            # Horizontal separator line + pivote text
            add_rect(slide, PAD_L, y, BOX_W, px(1), fill=BORDE_C)
            y += px(18 + 1)
            tf = add_tb(slide, PAD_L, y, px(760), H - y - px(20))
            add_multiline_element(tf, child, SYNE, Pt(16.5), NARANJA_OSC,
                                  space_between=Pt(6))
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.bold = True

    add_counter(slide, idx, total)


# ── Content-slide body sub-builders ──────────────────────────────────────────

def build_agenda_grid(slide, el, left, y, body_w):
    """4 agenda-cells in a row."""
    cells = el.select('.agenda-cell')
    n = len(cells)
    if n == 0:
        return y
    gap = px(12)
    cell_w = (body_w - gap * (n - 1)) // n
    cell_h = px(120)

    for i, cell in enumerate(cells):
        cx = left + i * (cell_w + gap)
        # Card background
        add_rect(slide, cx, y, cell_w, cell_h, fill=CARD_BG, line_color=BORDE_C, line_width=Pt(0.5))
        # Content
        tf = add_tb(slide, cx + px(10), y + px(12), cell_w - px(20), cell_h - px(20))
        mn = cell.select_one('.min')
        if mn:
            simple_para(tf, mn.get_text(strip=True), IBM_M, Pt(7.5), NARANJA_OSC,
                        space_after=Pt(3))
        tit = cell.select_one('.titulo')
        if tit:
            simple_para(tf, tit.get_text(strip=True), SYNE, Pt(10.5), AZUL,
                        bold=True, space_after=Pt(4))
        desc = cell.select_one('.desc')
        if desc:
            simple_para(tf, desc.get_text(strip=True), IBM_S, Pt(8.25), SEC,
                        space_after=Pt(0))
    return y + cell_h + px(12)


def build_two_col(slide, el, left, y, body_w):
    """2 col-cards side by side."""
    cards = el.select('.col-card')
    if not cards:
        return y
    gap = px(24)
    card_w = (body_w - gap) // 2
    card_h = px(150)

    for i, card in enumerate(cards):
        cx = left + i * (card_w + gap)
        add_rect(slide, cx, y, card_w, card_h, fill=CARD_BG, line_color=BORDE_C, line_width=Pt(0.5))
        tf = add_tb(slide, cx + px(12), y + px(12), card_w - px(24), card_h - px(20))

        lbl = card.select_one('.col-label')
        if lbl:
            cls = lbl.get('class', [])
            c = AZUL_MEDIO if 'label-azul' in cls else NARANJA_OSC
            simple_para(tf, lbl.get_text(strip=True), IBM_M, Pt(6.75), c,
                        space_after=Pt(6))

        titulo = card.select_one('.col-titulo')
        if titulo:
            simple_para(tf, titulo.get_text(strip=True), SYNE, Pt(12.75), AZUL,
                        bold=True, space_after=Pt(6))

        for p_tag in card.find_all('p'):
            p = tf.add_paragraph()
            p.space_before = Pt(0)
            p.space_after = Pt(3)
            for node in p_tag.children:
                parse_inline_node(node, p, IBM_S, Pt(9.75), TEXTO)

    return y + card_h + px(14)


def build_prda(slide, el, left, y, body_w):
    """PRDA boxes: P -> R -> D -> A in a flex row."""
    boxes = el.select('.prda-box')
    arrows = el.select('.prda-arrow')
    if not boxes:
        return y

    box_w = px(130)
    arrow_w = px(30)
    gap = px(8)
    total_w = len(boxes) * box_w + len(arrows) * arrow_w + (len(boxes) + len(arrows) - 1) * gap
    x_start = left + (body_w - total_w) // 2
    box_h = px(70)
    cx = x_start

    for i, box in enumerate(boxes):
        add_rect(slide, cx, y, box_w, box_h, fill=CARD_BG, line_color=AZUL, line_width=Pt(1.5))
        tf = add_tb(slide, cx, y + px(8), box_w, box_h - px(8))
        letra = box.select_one('.letra')
        if letra:
            simple_para(tf, letra.get_text(strip=True), SYNE, Pt(21), NARANJA_OSC,
                        bold=True, align=PP_ALIGN.CENTER, space_after=Pt(2))
        palabra = box.select_one('.palabra')
        if palabra:
            simple_para(tf, palabra.get_text(strip=True), SYNE, Pt(10.5), AZUL,
                        bold=True, align=PP_ALIGN.CENTER)
        cx += box_w + gap
        if i < len(arrows):
            tf2 = add_tb(slide, cx, y + px(24), arrow_w, px(30))
            simple_para(tf2, "->", SYNE, Pt(16.5), NARANJA_OSC,
                        bold=True, align=PP_ALIGN.CENTER)
            cx += arrow_w + gap

    return y + box_h + px(14)


def build_gate_grid(slide, el, left, y, body_w):
    """2 gate-cards side by side."""
    cards = el.select('.gate-card')
    if not cards:
        return y
    gap = px(18)
    card_w = (body_w - gap) // 2
    card_h = px(160)

    for i, card in enumerate(cards):
        cx = left + i * (card_w + gap)
        add_rect(slide, cx, y, card_w, card_h, fill=CARD_BG, line_color=BORDE_C, line_width=Pt(0.5))
        tf = add_tb(slide, cx + px(14), y + px(14), card_w - px(28), card_h - px(20))

        is_next = 'next' in card.get('class', [])
        lbl_color = NARANJA_OSC if is_next else VERDE
        bullet_color = lbl_color

        lbl = card.select_one('.gate-label')
        if lbl:
            simple_para(tf, lbl.get_text(strip=True), IBM_M, Pt(6.75), lbl_color,
                        space_after=Pt(8))

        for li in card.select('li'):
            p = tf.add_paragraph()
            p.space_before = Pt(0)
            p.space_after = Pt(4)
            r = p.add_run()
            r.text = "▸  "
            r.font.name = IBM_S
            r.font.size = Pt(9.75)
            r.font.color.rgb = bullet_color
            r.font.bold = True
            for node in li.children:
                parse_inline_node(node, p, IBM_S, Pt(9.75), TEXTO)

    return y + card_h + px(12)


def build_inline_grid_b5(slide, el, left, y, body_w):
    """Slide 13 custom grid: alternating label + description rows."""
    rows = list(el.children)
    # Group in pairs: label div + description div
    items = [c for c in rows if isinstance(c, Tag)]
    row_h = px(22)
    i = 0
    while i + 1 < len(items):
        label_el = items[i]
        desc_el = items[i + 1]
        col1_w = px(30)
        col2_w = body_w - col1_w - px(16)

        tf1 = add_tb(slide, left, y, col1_w, row_h)
        simple_para(tf1, label_el.get_text(strip=True), IBM_M, Pt(10.5), NARANJA_OSC,
                    bold=True)

        tf2 = add_tb(slide, left + col1_w + px(16), y, col2_w, row_h * 2)
        p = _first_para(tf2)
        p.space_after = Pt(2)
        for node in desc_el.children:
            parse_inline_node(node, p, IBM_S, Pt(10.5), TEXTO)
        y += row_h + px(4)
        i += 2
    return y


def build_content(slide, sec, html_dir, idx, total):
    """Navy header + crema body slides (agenda, content A-D, gate)."""
    set_slide_bg(slide, CREMA)

    # Header strip
    add_rect(slide, 0, 0, W, HDR_H, fill=AZUL)
    chip = sec.select_one('.content-chip')
    title_el = sec.select_one('.content-title')
    hdr_x = px(32)
    if chip:
        chip_txt = chip.get_text(strip=True)
        chip_tb_w = px(len(chip_txt) * 5 + 24)
        add_rect(slide, hdr_x, px(14), chip_tb_w, HDR_H - px(24),
                 fill=None, line_color=NARANJA_LUZ, line_width=Pt(0.5))
        tf = add_tb(slide, hdr_x + px(6), px(16), chip_tb_w - px(6), HDR_H - px(26))
        simple_para(tf, chip_txt, IBM_M, Pt(6), NARANJA_LUZ)
        hdr_x += chip_tb_w + px(16)
    if title_el:
        tf = add_tb(slide, hdr_x, px(14), W - hdr_x - px(32), HDR_H - px(24))
        simple_para(tf, title_el.get_text(strip=True), SYNE, Pt(12), WHITE, bold=True)

    # Body
    body = sec.select_one('.content-body')
    if not body:
        add_counter(slide, idx, total)
        return

    y = HDR_H + BODY_T
    left = BODY_L
    body_w = W - 2 * BODY_L

    for child in body.children:
        if not isinstance(child, Tag):
            continue
        tag = child.name.lower()
        cls = child.get('class', [])
        st = parse_inline_style(child.get('style', ''))

        if tag == 'h2':
            sz = css_px_to_pt(st.get('font-size', '')) or Pt(21)
            mt = css_px_to_emu(st.get('margin-top', ''))
            y += mt
            tf = add_tb(slide, left, y, body_w, px(50))
            add_multiline_element(tf, child, SYNE, sz, AZUL, space_between=Pt(2))
            for pp in tf.paragraphs:
                for r in pp.runs:
                    r.font.bold = True
            y += px(50)

        elif tag == 'p':
            txt = child.get_text(strip=True)
            if not txt:
                continue
            col = resolve_color(st.get('color')) or TEXTO
            sz = css_px_to_pt(st.get('font-size', '')) or Pt(11.25)
            mt = css_px_to_emu(st.get('margin-top', ''))
            y += mt
            lines_est = max(1, len(txt) // 90 + 1)
            h_est = px(18 * lines_est + 8)
            tf = add_tb(slide, left, y, body_w, h_est)
            p = _first_para(tf)
            p.space_after = Pt(6)
            for node in child.children:
                parse_inline_node(node, p, IBM_S, sz, col)
            y += h_est

        elif tag == 'ul':
            for li in child.find_all('li', recursive=False):
                li_lines = max(1, len(li.get_text(strip=True)) // 85 + 1)
                li_h = px(18 * li_lines + 6)
                tf = add_tb(slide, left, y, body_w, li_h)
                p = _first_para(tf)
                p.space_after = Pt(5)
                # bullet
                r0 = p.add_run()
                r0.text = "▸  "
                r0.font.name = IBM_S
                r0.font.size = Pt(11)
                r0.font.color.rgb = NARANJA_OSC
                r0.font.bold = True
                for node in li.children:
                    parse_inline_node(node, p, IBM_S, Pt(11.25), TEXTO)
                y += li_h

        elif 'agenda-grid' in cls:
            y = build_agenda_grid(slide, child, left, y, body_w)

        elif 'two-col' in cls:
            y = build_two_col(slide, child, left, y, body_w)

        elif 'prda' in cls:
            y = build_prda(slide, child, left, y, body_w)

        elif 'gate-grid' in cls:
            y = build_gate_grid(slide, child, left, y, body_w)

        elif tag == 'div':
            # Slide 13: inline CSS grid (B2-B6 session map)
            inner_st = parse_inline_style(child.get('style', ''))
            if 'grid-template-columns' in inner_st:
                y = build_inline_grid_b5(slide, child, left, y, body_w)
            else:
                txt = child.get_text(strip=True)
                if txt:
                    tf = add_tb(slide, left, y, body_w, px(50))
                    simple_para(tf, txt, IBM_S, Pt(11), TEXTO)
                    y += px(50)

    add_counter(slide, idx, total)


# ── Slide dispatcher ───────────────────────────────────────────────────────────

def detect_type(section):
    cls = section.get('class', [])
    if 'cover' in cls:
        return 'cover'
    if 'idea-slide' in cls:
        return 'idea'
    if 'image-slide' in cls:
        return 'image'
    if 'story' in cls:
        return 'story'
    return 'content'


def build_slide(prs, section, idx, total, html_dir):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    t = detect_type(section)
    if t == 'cover':
        build_cover(slide, section, idx, total)
    elif t == 'idea':
        build_idea(slide, section, idx, total)
    elif t == 'image':
        build_image_slide(slide, section, html_dir, idx, total)
    elif t == 'story':
        build_story(slide, section, idx, total)
    else:
        build_content(slide, section, html_dir, idx, total)


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    html_path = Path(sys.argv[1]).resolve()
    if not html_path.exists():
        print(f"ERROR: not found -- {html_path}")
        sys.exit(1)

    out_path = html_path.with_name(html_path.stem + "_editable.pptx")
    html_dir = html_path.parent

    print(f"Source : {html_path.name}")
    print(f"Output : {out_path.name}")

    soup = BeautifulSoup(html_path.read_text(encoding='utf-8'), 'html.parser')
    sections = soup.select('.slide')
    total = len(sections)
    print(f"Slides : {total}")

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    for i, sec in enumerate(sections, 1):
        print(f"  [{i:02d}/{total}] {detect_type(sec)}: {sec.get('class', [])}")
        build_slide(prs, sec, i, total, html_dir)

    prs.save(str(out_path))
    kb = out_path.stat().st_size // 1024
    print(f"Done -> {out_path.name}  ({kb} KB)")


if __name__ == "__main__":
    main()
