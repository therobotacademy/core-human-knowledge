#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_docx.py — PROTOTIPO de repo-code-explainer · etapas 3-4.

3. RASTERIZAR: <OUT>/svg/*.svg → <OUT>/png/*.png  (PyMuPDF / fitz, DPI 150).
4. ENSAMBLAR: DOCX con portada + una sección por unidad (narrativa + imagen embebida).

Embebe cada diagrama como **par PNG + SVG** (capa vectorial vía `asvg:svgBlip`, el mismo
patrón del precedente II3): Word renderiza el SVG nítido y cae al PNG en visores antiguos.
Sin dependencia de Node — se hace con `python-docx` + inyección OOXML (ver `attach_svg`).

    python .claude/skills/repo-code-explainer/build_docx.py
"""
from pathlib import Path
import fitz

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO = Path(__file__).resolve().parents[3]
OUT = REPO / "CONTENT" / "PARTE2-coiiaoc-Materiales" / "II-CrewIA" / "M_codigo-crew_NCR"
SVG_DIR, PNG_DIR = OUT / "svg", OUT / "png"
DOCX = OUT / "M_codigo_MVP-Crew_NCR.docx"

# Paleta v1.1
AZUL, NAR_OSC, CREMA, TEXSEC = RGBColor(0x1E, 0x3A, 0x5F), RGBColor(0xC2, 0x51, 0x0A), RGBColor(0xFA, 0xFA, 0xF7), RGBColor(0x6B, 0x6B, 0x6B)

# Unidad → (título, narrativa "por qué")
UNITS = [
    ("01-orquestacion-crew", "Orquestación del Crew (main.py)",
     "El punto de entrada arma las 3 Task con el texto del NCR inyectado y las ejecuta en un "
     "Crew secuencial (Process.sequential). La salida se reconstruye mapeando cada "
     "tasks_output[i].pydantic a su tramo del JSON final. El verbose=True es deliberado: en clase "
     "se ve el razonamiento de los tres agentes."),
    ("02-agente-clasificador", "Agente · clasificador (AS9100)",
     "Configuración declarativa: role/goal/backstory fijan el comportamiento; la tool lookup_tipo_nc "
     "ancla la severidad y la cláusula al catálogo del sistema (no a la opinión del LLM). El backstory "
     "fuerza SIEMPRE confirmar con el catálogo — un patrón anti-alucinación por diseño de prompt."),
    ("03-agente-causas", "Agente · causas (análisis 8D)",
     "Segundo eslabón: realiza el 8D D1-D5. Igual que el clasificador, delega la causa raíz en un "
     "catálogo (lookup_causas_raiz) parametrizado por el área 8D, y solo refina la redacción con el "
     "detalle del NCR. La consistencia con la base de datos manda sobre la creatividad."),
    ("04-agente-planes", "Agente · planes (CAPA)",
     "Tercer eslabón, SIN tool: es razonamiento puro sobre el 8D recibido. Genera el plan CAPA "
     "(mínimo 2 acciones: inmediata + preventiva), responsables por área (no personas, por privacidad) "
     "y una confianza global que señala cuándo el NCR es incompleto."),
    ("05-cadena-tasks", "Cadena de Tasks (tasks.py)",
     "El acoplamiento entre agentes NO es por variable en memoria: cada Task declara context=[anterior] "
     "y output_pydantic=Modelo. CrewAI encadena las salidas y el provider (Anthropic) las devuelve como "
     "structured output validado. Es el mismo principio de acoplamiento-por-contrato que un pipeline."),
    ("06-modelos-pydantic", "Modelos Pydantic (contrato de salida)",
     "Los cuatro BaseModel son el contrato que el LLM debe cumplir. Literal[...] restringe la severidad "
     "a 3 valores; Field(gt=0) y Field(ge=0, le=1) validan plazos y confianza. Definir el esquema aquí "
     "es lo que convierte texto libre del LLM en datos fiables aguas abajo."),
    ("07-tool-lookup-tipo-nc", "Tool · lookup_tipo_nc",
     "Bifurcación con red de seguridad: intenta match exacto, luego parcial (substring), y si nada "
     "encaja devuelve un fallback conservador (Major · 8.7 · contención) con un aviso. La lógica vive en "
     "_lookup_tipo_nc_impl (testeable) separada del wrapper @tool que usa el agente."),
    ("08-tool-lookup-causas", "Tool · lookup_causas_raiz",
     "Valida primero el área contra un conjunto cerrado (_AREAS_VALIDAS) y corta con error si no es "
     "válida — un guardarraíl antes de tocar el catálogo. Luego filtra, ordena por frecuencia "
     "(Alta→Media→Baja) y devuelve el top-3. Determinista y auditable."),
]


def _p(doc, text, size=11, bold=False, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "IBM Plex Sans"
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def rasterize():
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    for svg in sorted(SVG_DIR.glob("*.svg")):
        doc = fitz.open(str(svg))
        pdf = fitz.open("pdf", doc.convert_to_pdf())
        png = PNG_DIR / (svg.stem + ".png")
        pdf[0].get_pixmap(dpi=150).save(str(png))
    return len(list(PNG_DIR.glob("*.png")))


def attach_svg(doc, shape, svg_path):
    """Adjunta la capa vectorial SVG al blip PNG (patrón Word PNG+SVG · ext asvg:svgBlip, como II3).

    Word muestra el SVG (nítido al escalar) y cae al PNG en visores antiguos.
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.part import Part
    from docx.oxml import parse_xml
    pkg = doc.part.package
    partname = pkg.next_partname("/word/media/image%d.svg")
    svg_part = Part(partname, "image/svg+xml", Path(svg_path).read_bytes(), pkg)
    rid = doc.part.relate_to(svg_part, RT.IMAGE)
    blip = shape._inline.xpath(".//a:blip")[0]
    blip.append(parse_xml(
        '<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
        '<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:embed="%s"/></a:ext></a:extLst>' % rid))


def build():
    doc = Document()
    # portada
    _p(doc, "MVP-Crew_NCR", size=26, bold=True, color=AZUL, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    _p(doc, "Documentación del código · Crew CrewAI para análisis de No Conformidades (NCR)",
       size=13, color=NAR_OSC, space_after=2)
    _p(doc, "Generado con repo-code-explainer (prototipo) · paleta COIIAOC v1.1 · 8 unidades",
       size=10, color=TEXSEC, space_after=14)
    _p(doc, "Arquitectura: 3 agentes (clasificador → causas → planes) encadenados en un Crew secuencial, "
            "cada uno con structured output Pydantic; 2 tools de catálogo (AS9100 / 8D) anclan los valores "
            "al sistema. Cada sección explica una unidad de código con su diagrama.", size=11, space_after=10)

    for stem, title, why in UNITS:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        _p(doc, title, size=15, bold=True, color=AZUL, space_after=4)
        _p(doc, why, size=11, space_after=8)
        png = PNG_DIR / (stem + ".png")
        svg = SVG_DIR / (stem + ".svg")
        if png.exists():
            shape = doc.add_picture(str(png), width=Inches(6.3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if svg.exists():
                attach_svg(doc, shape, svg)   # capa vectorial (PNG + SVG, patrón II3)

    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DOCX))
    return DOCX


def main():
    n = rasterize()
    print(f"[OK] rasterizados {n} PNG → {PNG_DIR.relative_to(REPO)}")
    out = build()
    print(f"[OK] DOCX → {out.relative_to(REPO)}")


if __name__ == "__main__":
    main()
